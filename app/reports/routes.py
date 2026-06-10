# app/reports/routes.py

from flask import render_template, request, session, redirect, url_for, send_file, flash
from flask_login import login_required, current_user
from app.reports import bp
from app import db
from app.models import Asset, Personnel, MaintenanceLog, UploadHistory
from app.utils import to_upper, format_phone, normalize_text 
from datetime import datetime
import io
import qrcode
import base64
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import func, or_

# --- 1. DOWNLOAD EXCEL TEMPLATE ---
@bp.route('/download-template/<type>')
@login_required
def download_template(type):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sample Template"
    bold_font = Font(bold=True)
    
    if type == 'personnel':
        ws.append(['Full Name', 'Title', 'Department', 'Campus', 'Office', 'Phone', 'Email'])
        ws.append(['John Doe', 'Clerk', 'IT', 'Main Campus', 'Z-23', '5551234567', 'john@example.com'])
        ws.column_dimensions['A'].width = 25
    else:
        # Asset Template
        ws.append(['Asset Name', 'Type', 'Brand', 'Model', 'Serial No', 'Campus', 'Location', 'Quantity', 'Assigned To', 'Description'])
        ws.append(['Work Desk', 'Furniture', '', '', '', 'Main Campus', 'B-Block 105', 1, '', ''])
        ws.column_dimensions['A'].width = 25

    for cell in ws[1]: cell.font = bold_font
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=f"Sample_{type.capitalize()}_Template.xlsx", mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# --- 2. GENERAL REPORT (Download Entire DB) ---
@bp.route('/report')
@login_required
def report():
    if int(session.get('auth_level', 0)) < 1:
        return redirect(url_for('main.index'))

    wb = Workbook()
    
    # SHEET 1: ASSETS
    ws1 = wb.active
    ws1.title = "Asset List"
    ws1.append(['ID', 'Asset Name', 'Type', 'Campus', 'Location', 'Quantity', 'Purchase Date'])
    
    query_d = Asset.query.all()
    
    for d in query_d:
        ws1.append([d.id, d.name, d.type, d.campus, d.location, d.quantity, d.date_added])

    # SHEET 2: PERSONNEL
    ws2 = wb.create_sheet("Personnel List")
    ws2.append(['ID', 'Full Name', 'Title', 'Department', 'Campus', 'Office', 'Phone', 'Email'])
    
    for p in Personnel.query.all():
        ws2.append([p.id, p.full_name, p.title, p.department, p.campus, p.office, p.phone, p.email])

    # SHEET 3: MAINTENANCE
    ws3 = wb.create_sheet("Maintenance Logs")
    ws3.append(['ID', 'Location', 'Title', 'Status', 'Reported By', 'Date'])
    
    for a in MaintenanceLog.query.all():
        # User query to find reporter is not directly mapped as user relation might be null
        reporter = str(a.user_id)
        ws3.append([a.id, a.location, a.title, a.status, reporter, a.date_reported])

    # Style Settings
    bold = Font(bold=True)
    for sheet in wb.worksheets:
        for cell in sheet[1]: cell.font = bold
        sheet.column_dimensions['A'].width = 10
        sheet.column_dimensions['B'].width = 25

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=f"General_Report_{datetime.now().strftime('%Y-%m-%d')}.xlsx")

# --- 3. BULK PRINT QR CODES ---
@bp.route('/bulk-print-assets', methods=['GET', 'POST'])
@login_required
def bulk_print_assets():
    if request.method == 'POST': 
        selected_ids = request.form.getlist('selected_ids')
    else: 
        selected_ids = request.args.get('ids', '').split(',')
    
    if not selected_ids or selected_ids == ['']: 
        return redirect(url_for('main.index', tab='assets'))
    
    assets = Asset.query.filter(Asset.id.in_(selected_ids)).all()
    qr_list = []
    
    for item in assets:
        date_info = item.date_added if item.date_added else "Not Specified"
        qr_content = f"ASSET INFO\nID: {item.id}\nItem: {item.name}\nLocation: {item.location}\nType: {item.type}\nDate: {date_info}"
        
        qr = qrcode.QRCode(box_size=10, border=2)
        qr.add_data(qr_content)
        qr.make(fit=True)
        img = qr.make_image(fill_color="black", back_color="white")
        
        buffered = io.BytesIO()
        img.save(buffered, format="PNG")
        qr_base64 = base64.b64encode(buffered.getvalue()).decode('utf-8')
        
        qr_list.append({
            'name': item.name, 'id': item.id, 'location': item.location,
            'type': item.type, 'qr': qr_base64
        })

    return render_template('toplu_yazdir.html', qr_list=qr_list)

# --- 4. DOOR CARD & SINGLE QR ---
@bp.route('/door-card/<path:location_name>')
def door_card(location_name):
    target_url = url_for('reports.office_details', location_name=location_name, _external=True)
    
    qr = qrcode.QRCode(box_size=10, border=2)
    qr.add_data(target_url)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    
    buffered = io.BytesIO()
    img.save(buffered, format="PNG")
    qr_base64 = base64.b64encode(buffered.getvalue()).decode('utf-8')
    
    return render_template('kapi_karti.html', location=location_name, qr_code=qr_base64)

@bp.route('/office/<path:location_name>')
def office_details(location_name):
    t = f"%{normalize_text(location_name)}%"
    
    assets = Asset.query.filter(func.NORMALIZE(Asset.location).like(t)).all()
    personnel = Personnel.query.filter(func.NORMALIZE(Personnel.office).like(t)).all()
    
    return render_template('oda_detay.html', location=location_name, assets=assets, personnel=personnel)

@bp.route('/report-analysis')
@login_required
def report_analysis():
    if int(session.get('auth_level', 0)) < 1:
        return redirect(url_for('main.index'))

    analysis_type = request.args.get('analysis_type', 'assets')
    
    req_item = request.args.get('req_item', '').strip()
    req_campus = request.args.get('req_campus', '')
    req_location = request.args.get('req_location', '').strip()
    
    req_p_name = request.args.get('req_p_name', '').strip()
    req_p_dept = request.args.get('req_p_dept', '').strip()
    req_p_campus = request.args.get('req_p_campus', '')
    req_p_office = request.args.get('req_p_office', '').strip()

    wb = Workbook()
    ws = wb.active
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="198754", fill_type="solid")
    grand_total = 0

    if analysis_type == 'personnel':
        ws.title = "Personnel Analysis"
        ws.append(['No', 'Full Name', 'Title', 'Department', 'Campus', 'Office', 'Phone', 'Email'])
        
        q = Personnel.query
        if req_p_name:
            t = f"%{normalize_text(req_p_name)}%"
            q = q.filter(or_(func.NORMALIZE(Personnel.full_name).like(t), func.NORMALIZE(Personnel.title).like(t)))
        if req_p_dept:
            q = q.filter(func.NORMALIZE(Personnel.department).like(f"%{normalize_text(req_p_dept)}%"))
        if req_p_campus and req_p_campus != "All":
            q = q.filter(Personnel.campus == req_p_campus)
        if req_p_office:
            q = q.filter(func.NORMALIZE(Personnel.office).like(f"%{normalize_text(req_p_office)}%"))
            
        results = q.order_by(Personnel.department.asc(), Personnel.full_name.asc()).all()
        
        for i, p in enumerate(results, 1):
            phone_fmt = format_phone(p.phone)
            ws.append([i, p.full_name, p.title, p.department, p.campus, p.office, phone_fmt, p.email])
            grand_total += 1
            
        ws.append(['', '', '', '', '', '', 'GRAND TOTAL:', grand_total])

    else:
        ws.title = "Asset Analysis"
        ws.append(['No', 'Asset Name', 'Campus', 'Location / Office', 'Quantity'])
        
        q = db.session.query(Asset.name, Asset.campus, Asset.location, func.sum(Asset.quantity).label('total_qty'))
        
        if req_item:
            q = q.filter(func.NORMALIZE(Asset.name).like(f"%{normalize_text(req_item)}%"))
        if req_campus and req_campus != "All":
            q = q.filter(Asset.campus == req_campus)
        if req_location:
            q = q.filter(func.NORMALIZE(Asset.location).like(f"%{normalize_text(req_location)}%"))
            
        results = q.group_by(Asset.name, Asset.campus, Asset.location)\
                    .order_by(Asset.campus.asc(), Asset.location.asc()).all()
                    
        for i, item in enumerate(results, 1):
            ws.append([i, item.name, item.campus, item.location, item.total_qty])
            grand_total += item.total_qty
            
        ws.append(['', '', '', 'GRAND TOTAL:', grand_total])

    # Style Settings
    for cell in ws[1]: 
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 30

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    file_name = f"Analysis_Report_{analysis_type}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(output, download_name=file_name, as_attachment=True)

@bp.route('/report-chart-summary')
@login_required
def report_chart_summary():
    analysis_type = request.args.get('analysis_type', 'assets')
    
    req_item = request.args.get('req_item', '').strip()
    req_campus = request.args.get('req_campus', '')
    req_p_dept = request.args.get('req_p_dept', '').strip()

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Campus Distribution"
    
    ws1.append(['Campus Name', 'Count'])
    
    temp_campus = {}
    
    if analysis_type == 'personnel':
        q = Personnel.query
        if req_p_dept:
            q = q.filter(func.NORMALIZE(Personnel.department).like(f"%{normalize_text(req_p_dept)}%"))
        
        for p in q.all():
            k = p.campus if p.campus else "Not Specified"
            temp_campus[k] = temp_campus.get(k, 0) + 1
    else:
        q = Asset.query
        if req_item:
            q = q.filter(func.NORMALIZE(Asset.name).like(f"%{normalize_text(req_item)}%"))
        if req_campus and req_campus != "All":
            q = q.filter(Asset.campus == req_campus)
            
        for d in q.all():
            temp_campus[d.campus] = temp_campus.get(d.campus, 0) + d.quantity

    for k, v in temp_campus.items():
        ws1.append([k, v])
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(output, download_name="Chart_Summary_Data.xlsx", as_attachment=True)