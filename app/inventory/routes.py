from flask import redirect, url_for, request, session, flash, jsonify, render_template
from app.inventory import bp
from app import db
from app.models import Asset, UploadHistory, MaintenanceLog
from datetime import datetime
from flask_login import login_required, current_user
from app.utils import normalize_text, to_upper, to_title, save_log
from sqlalchemy import text
import openpyxl
import os
import uuid


# --- CRUD OPERATIONS ---

@bp.route('/add-asset', methods=['POST'])
@login_required
def add_asset():
    if session.get('role') == 'technician':
        return redirect(url_for('main.index', tab='maintenance'))
    
    name = request.form.get('name')
    type_ = request.form.get('type')
    campus = request.form.get('campus')
    location = request.form.get('location')
    quantity = request.form.get('quantity')
    
    new_asset = Asset(
        name=name,
        type=type_,
        campus=campus,
        location=location,
        quantity=quantity,
        date_added=datetime.now().strftime("%Y-%m-%d"),
        asset_tag=str(uuid.uuid4())[:8].upper()
    )
    
    db.session.add(new_asset)
    db.session.commit()
    
    save_log(f"{name} Added", f"Location: {location}", "Add")
    flash(f"{name} added successfully.", "success")
    return redirect(url_for('main.index', tab='assets'))

@bp.route('/update-asset', methods=['POST'])
@login_required
def update_asset():
    if int(session.get('auth_level', 0)) < 1:
        flash("Unauthorized action.", "danger")
        return redirect(url_for('main.index', tab='assets'))
    d_id = request.form.get('id')
    
    asset = Asset.query.get(d_id)
    
    if asset:
        asset.name = request.form.get('name')
        asset.type = request.form.get('type')
        asset.campus = request.form.get('campus')
        asset.location = request.form.get('location')
        asset.quantity = request.form.get('quantity')
        
        db.session.commit()
        save_log(f"{asset.name} Updated", f"New Location: {asset.location}", "Edit")
        flash("Record updated.", "success")
        
    return redirect(url_for('main.index', tab='assets'))

@bp.route('/move-asset', methods=['POST'])
@login_required
def move_asset():
    if int(session.get('auth_level', 0)) < 1:
        flash("Unauthorized action.", "danger")
        return redirect(url_for('main.index', tab='assets'))

    d_id = request.form.get('id')
    new_campus = request.form.get('campus')
    new_location = request.form.get('location')
    
    asset = Asset.query.get(d_id)
    
    if asset:
        old_location = f"{asset.campus}/{asset.location}"
        asset.campus = new_campus
        asset.location = new_location
        
        db.session.commit()
        save_log(f"{asset.name} Moved", f"{old_location} -> {new_campus}/{new_location}", "Move")
        flash("Asset moved successfully.", "success")
        
    return redirect(url_for('main.index', tab='assets'))

@bp.route('/delete-asset/<int:id>')
@login_required
def delete_asset(id):
    if int(session.get('auth_level', 0)) < 3:
        flash("No permission to delete.", "danger")
        return redirect(url_for('main.index', tab='assets'))
        
    asset = Asset.query.get(id)
    if asset:
        save_log(f"{asset.name} Deleted", f"Old Location: {asset.location}", "Delete")
        db.session.delete(asset)
        db.session.commit()
        flash("Asset deleted.", "warning")
        
    return redirect(url_for('main.index', tab='assets'))

@bp.route('/reset-assets')
@login_required
def reset_assets():
    if int(session.get('auth_level', 0)) < 3:
        return redirect(url_for('main.index'))
        
    try:
        db.session.query(Asset).delete()
        if db.engine.name == 'sqlite':
            try:
                db.session.execute(text("DELETE FROM sqlite_sequence WHERE name='asset'"))
            except Exception:
                pass
        db.session.commit()
        
        save_log("All List Deleted", "Database Reset", "Reset")
        flash("All asset list cleared.", "danger")
    except Exception as e:
        db.session.rollback()
        flash(f"Error: {str(e)}", "danger")
        
    return redirect(url_for('main.index', tab='assets'))

# --- BULK OPERATIONS ---

@bp.route('/bulk-move-assets', methods=['POST'])
@login_required
def bulk_move_assets():
    if session.get('auth_level', 0) < 1 and session.get('role') != 'technician':
        return jsonify({'status': 'error', 'msg': 'Unauthorized action!'})

    selected_ids = request.form.getlist('selected_ids')
    new_campus = request.form.get('new_campus')
    new_location = request.form.get('new_location')

    if not selected_ids:
        flash("No selection made.", "warning")
        return redirect(url_for('main.index', tab='assets'))

    query = Asset.query.filter(Asset.id.in_(selected_ids))
    
    count = 0
    for item in query.all():
        if new_campus: item.campus = new_campus
        if new_location: item.location = new_location
        count += 1
        
    db.session.commit()
    flash(f"{count} assets moved.", "success")
    return redirect(url_for('main.index', tab='assets'))

@bp.route('/bulk-delete-assets', methods=['POST'])
@login_required
def bulk_delete_assets():
    if int(session.get('auth_level', 0)) < 3:
        return redirect(url_for('main.index', tab='assets'))

    selected_ids = request.form.getlist('selected_ids')
    if not selected_ids: return redirect(url_for('main.index', tab='assets'))
    
    deleted_count = Asset.query.filter(Asset.id.in_(selected_ids)).delete(synchronize_session=False)
    db.session.commit()
    
    save_log(f"{deleted_count} Records Deleted", "Bulk Action", "Delete")
    flash(f"{deleted_count} records deleted.", "success")
    
    return redirect(url_for('main.index', tab='assets'))

@bp.route('/upload-assets', methods=['POST'])
@login_required
def upload_assets():
    if 'file' not in request.files: 
        return redirect(url_for('main.index'))
        
    files = request.files.getlist('file')
    target_campus = request.form.get('target_campus', 'Main Campus')
    building_floor = request.form.get('building_floor', '')
    
    total_added = 0
    processed = False

    for file in files:
        if file.filename == '': continue
        try:
            file_name_clean = file.filename.rsplit('.', 1)[0].replace('_', ' ').title()
            
            wb = openpyxl.load_workbook(file)
            for ws in wb.worksheets:
                full_location = " / ".join([p for p in [building_floor, file_name_clean, ws.title] if p])
                
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or row[0] is None: continue
                    
                    name = row[0]
                    type_ = row[1] if len(row) > 1 else ""
                    
                    try: 
                        quantity = int(row[2]) if len(row) > 2 and row[2] else 1
                    except: 
                        quantity = 1
                    
                    existing = Asset.query.filter_by(name=name, location=full_location).first()
                    
                    if not existing:
                        new_item = Asset(
                            name=name,
                            type=type_,
                            campus=target_campus,
                            location=full_location,
                            quantity=quantity,
                            date_added=datetime.now().strftime("%Y-%m-%d"),
                            asset_tag=str(uuid.uuid4())[:8].upper()
                        )
                        db.session.add(new_item)
                        total_added += 1
                        processed = True

        except Exception as e:
            flash(f"Error ({file.filename}): {str(e)}", "danger")
            continue

    if processed:
        db.session.commit()
        save_log(f"{len(files)} Files Uploaded", f"{building_floor} - {total_added} Assets", "Upload")
        flash(f"{total_added} assets added to the system.", "success")
    else:
        flash("No new assets added or all already registered.", "warning")

    return redirect(url_for('main.index', tab='assets'))


@bp.route('/upload-folder', methods=['POST'])
@login_required
def upload_folder():
    if 'file' not in request.files: return redirect(url_for('main.index'))
    files = request.files.getlist('file')
    
    default_campus = request.form.get('target_campus', 'Main Campus')
    
    process_count = 0
    saved_locations = set()

    CAMPUS_MAP = {
        "pelitli": "North Campus",
        "çimenli": "South Campus",
        "cimenli": "South Campus",
        "kaşüstü": "East Campus",
        "kasustu": "East Campus",
        "yomra": "West Campus",
        "yalıncak": "Main Campus",
        "yalincak": "Main Campus",
        "ömer yıldız": "Main Campus",
        "omer yildiz": "Main Campus"
    }

    WORDS_TO_REMOVE = [
        "LİSTESİ", "LISTESI", "LİSTE", "LISTE", "DEMİRBAŞLARI", "DEMİRBAŞ", "DEMIRBAS", 
        "ENVANTER", "SAYIM", "SİSTEMİ", "SISTEMI", "YAPILDI", "YAPILAN", "YENİ", "ESKİ", 
        "COPY", "KOPYA", "YEDEK", "REVİZE", "REVIZE", "DÜZENLEME", "DÜZENLENEN", 
        "KONTROL", "TASLAK", "SON", "FİNAL", "FINAL", "MASAÜSTÜ", "DOWNLOADS", 
        "BELGELERİM", "TABLO", "TÜMÜ", "TUMU", "XLSX", "XLS",
        "YALINCAK", "PELİTLİ", "PELITLI", "ÇİMENLİ", "CIMENLI", "KAŞÜSTÜ", "KASUSTU", "YOMRA",
        "KANUNİ", "MERKEZ", "KAMPÜSÜ", "KAMPUSU", "YERLEŞKESİ", "YERLESKESI",
        "ÖMER YILDIZ", "OMER YILDIZ", "ÖMER", "YILDIZ"
    ]

    IMPORTANT_WORDS = [
        "BLOK", "KAT", "ODA", "BİNA", "BINA", "YURT", "OFİS", "OFFICE", 
        "HALL", "SALON", "LAB", "DEPO", "ZEMİN", "GİRİŞ", "SİSTEM", "KAZAN",
        "RESTORAN", "YEMEKHANE", "KANTİN", "LOBİ", "MESCİT", "GUVENLIK", "GÜVENLİK",
        "AMBAR", "ATÖLYE", "ARŞİV", "LİSE", "FAKÜLTE", "MYO", "MEMUR", "PERSONEL",
        "PATOLOJİ", "KLİNİK", "POLİKLİNİK", "SERVİS", "BÖLÜM", "BOLUM", "BİRİM", "LABORATUVAR"
    ]

    for file in files:
        if not (file.filename.endswith('.xlsx') or file.filename.endswith('.xls')):
            continue
        if '~$' in file.filename: continue
            
        try:
            full_path = file.filename.replace('\\', '/')
            path_lower = full_path.lower()
            
            active_campus = default_campus
            for key, real_name in CAMPUS_MAP.items():
                if key in path_lower:
                    active_campus = real_name
                    break 
            
            path_parts = full_path.split('/')
            file_name_raw = path_parts[-1].rsplit('.', 1)[0]
            all_parts = path_parts[:-1] + [file_name_raw]
            
            meaningful_path_parts = []

            for part in all_parts:
                clean_part = tr_upper(part)
                
                for banned in WORDS_TO_REMOVE:
                    if banned in clean_part:
                        clean_part = clean_part.replace(banned, "")
                
                clean_part = clean_part.replace("_", " ").replace("-", " ").strip()
                
                if len(clean_part) < 2 and not any(c.isdigit() for c in clean_part):
                    continue

                is_important = any(k in clean_part for k in IMPORTANT_WORDS)
                is_block_code = (len(clean_part) > 0 and len(clean_part) < 6 and any(c.isdigit() for c in clean_part))
                
                if (is_important or is_block_code or len(clean_part) > 2):
                    clean_part_title = tr_title(clean_part)
                    if meaningful_path_parts:
                        last_added = meaningful_path_parts[-1]
                        if clean_part_title in last_added or last_added in clean_part_title:
                            if len(clean_part_title) > len(last_added):
                                meaningful_path_parts[-1] = clean_part_title
                            continue 
                    
                    meaningful_path_parts.append(clean_part_title)

            clean_path_str = " / ".join(meaningful_path_parts)

            wb = openpyxl.load_workbook(file)
            
            for ws in wb.worksheets:
                sheet_name = ws.title.strip()
                
                if "Sheet" in sheet_name or "Sayfa" in sheet_name:
                    full_location = clean_path_str if clean_path_str else "General Storage"
                else:
                    if clean_path_str:
                        if sheet_name in clean_path_str:
                            full_location = clean_path_str
                        else:
                            full_location = f"{clean_path_str} / {sheet_name}"
                    else:
                        full_location = sheet_name 

                if len(full_location) > 150: full_location = full_location[:147] + "..."
                saved_locations.add(full_location)

                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or row[0] is None: continue
                    name = row[0]
                    type_ = row[1] if len(row)>1 else ""
                    try: quantity = int(row[2]) if len(row)>2 else 1
                    except: quantity = 1
                    
                    existing = Asset.query.filter_by(name=name, location=full_location).first()
                    
                    if not existing:
                        new_item = Asset(
                            name=name, 
                            type=type_, 
                            campus=active_campus, 
                            location=full_location, 
                            quantity=quantity, 
                            date_added=datetime.now().strftime("%Y-%m-%d"),
                            asset_tag=str(uuid.uuid4())[:8].upper()
                        )
                        db.session.add(new_item)
            
            process_count += 1

        except Exception as e:
            print(f"Error ({file.filename}): {e}")

    db.session.commit() 
    
    if process_count > 0:
        sample_location = list(saved_locations)[0] if len(saved_locations) > 0 else ""
        save_log(f"Smart Upload ({process_count} files)", f"Ex: {sample_location}", "Upload")
        flash(f"{process_count} files processed successfully.", "success")
    else:
        flash("No files found to process.", "warning")

    return redirect(url_for('main.index', tab='assets'))

@bp.route('/asset-details/<int:id>')
def asset_details(id):
    asset = Asset.query.get_or_404(id)
    
    active_maintenance = MaintenanceLog.query.filter(
        MaintenanceLog.asset_id == id,
        ~MaintenanceLog.status.in_(['Completed', 'Cancelled'])
    ).order_by(MaintenanceLog.id.desc()).first()
    
    return render_template('detay.html', item=asset, ariza=active_maintenance)